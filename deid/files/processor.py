"""Writes de-identified copies of CSV and Excel files (streaming CSV, Excel in place or streaming),
with a value cache and parallel pre-masking for large files."""

from __future__ import annotations

import csv
import datetime as _dt
import hashlib
import json
import logging
import os
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ..config import Rules
from ..masking import BaseMasker, KeyRing, MaskerError, build_maskers, column_masker
from .rules import FileColumn, FileRule, FileRuleError, FileRules, file_format, output_path_for

log = logging.getLogger("deid.files")


_MISSING = object()


CACHE_LIMIT = 2000000          # distinct (column, value) masks kept in memory (~300 MB worst case)


PARALLEL_MIN_BYTES = 5 * 1024 * 1024   # smaller files are faster without starting worker processes


PARALLEL_BATCH = 5000


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


_INT_RE = re.compile(r"^-?\d+$")


def cell_to_text(value: Any) -> Optional[str]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, _dt.datetime):
        return value.isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def text_to_cell(original: Any, masked: Optional[str]) -> Any:
    """Write back with the original cell type where possible (numbers stay numbers, dates stay dates)."""
    if masked is None:
        return None
    if isinstance(original, _dt.datetime):
        try:
            return _dt.datetime.fromisoformat(masked)
        except ValueError:
            return masked
    if isinstance(original, _dt.date):
        try:
            return _dt.date.fromisoformat(masked[:10])
        except ValueError:
            return masked
    if isinstance(original, (int, float)) and not isinstance(original, bool):
        if _INT_RE.match(masked or ""):
            return int(masked)
        try:
            return float(masked)
        except ValueError:
            return masked
    return masked


@dataclass
class FileResult:
    input_path: str
    output_path: str
    rule: str
    format: str
    status: str = "PENDING"
    rows: int = 0
    sheets: List[str] = field(default_factory=list)
    masked_cells: Dict[str, int] = field(default_factory=dict)
    value_errors: int = 0
    formulas_skipped: int = 0
    input_sha256: str = ""
    output_sha256: str = ""
    seconds: float = 0.0
    message: str = ""


class FileDeidentifier:
    def __init__(self, rules: Rules, file_rules: FileRules, keyring: KeyRing, workers: int = 1):
        self.rules = rules
        self.file_rules = file_rules
        self.keyring = keyring
        self.workers = max(1, int(workers))
        self._pool = None
        self.maskers = build_maskers(rules.domains.values(), keyring)
        self._col_maskers: Dict[Tuple[str, str], BaseMasker] = {}
        self._cache: Dict[Tuple[int, Optional[str], str], Optional[str]] = {}
        self._cache_ids: Dict[Tuple[str, str], int] = {}  # one small int per distinct (domain, overrides)
        self._col_ids: Dict[int, int] = {}

    def _masker(self, rule: FileRule, col: FileColumn) -> BaseMasker:
        key = (rule.name, col.name)
        m = self._col_maskers.get(key)
        if m is None:
            m = column_masker(self.maskers[col.domain], self.keyring, col.overrides)
            self._col_maskers[key] = m
        return m

    def _column_cache_id(self, col: FileColumn) -> int:
        """Columns with the same domain and format overrides share cached masks (they mask identically)."""
        cid = self._col_ids.get(id(col))
        if cid is None:
            sig = (col.domain, json.dumps(col.overrides, sort_keys=True) if col.overrides else "")
            cid = self._cache_ids.setdefault(sig, len(self._cache_ids))
            self._col_ids[id(col)] = cid
        return cid

    def _cache_key(self, col: FileColumn, text: str, entity: Optional[str]):
        return (self._column_cache_id(col), entity if col.entity_column else None, text)

    def close(self) -> None:
        if self._pool is not None:
            self._pool.close()
            self._pool = None

    def prewarm(self, rule: FileRule, targets, rows, to_text) -> int:
        """Mask the distinct values of a large file in parallel worker processes and fill the cache.

        Identical results to in-process masking (same maskers rebuilt from the same domains and key).
        Values that fail are left out of the cache so the main pass raises the usual error for them.
        """
        from ..engine.parallel import MaskingPool
        pending: Dict[Tuple[str, str], Dict[Any, Tuple[str, Optional[str]]]] = {}
        for values in rows:
            for col, pos, ent in targets:
                if pos >= len(values):
                    continue
                text = to_text(values[pos])
                if text is None or text.strip() == "":
                    continue
                entity = to_text(values[ent]) if ent is not None and ent < len(values) else None
                ck = self._cache_key(col, text, entity)
                if ck not in self._cache:
                    pending.setdefault((col.name, col.domain), {})[ck] = (text, entity)
            if sum(len(v) for v in pending.values()) + len(self._cache) >= CACHE_LIMIT:
                break
        if not pending:
            return 0
        if self._pool is None:
            self._pool = MaskingPool(self.workers, self.rules.domains.values(), self.keyring)
        done = 0
        for (col_name, _), items in pending.items():
            col = next(c for c in rule.columns if c.name == col_name)
            keys = list(items)
            # row ids travel as 8-byte keys (mask_rows passes keys through bytes())
            batches = [[(i.to_bytes(8, "big"), items[k][0], items[k][1])
                        for i, k in enumerate(keys[s:s + PARALLEL_BATCH], start=s)]
                       for s in range(0, len(keys), PARALLEL_BATCH)]
            for _, results in self._pool.map(self._masker(rule, col), col.overrides, batches, False, parallel=True):
                for idx, masked, err, _ in results:
                    if err is None:
                        self._cache[keys[int.from_bytes(idx, "big")]] = masked
                        done += 1
        return done

    def mask_text(self, rule: FileRule, col: FileColumn, text: Optional[str], entity: Optional[str]) -> Optional[str]:
        if text is None or text.strip() == "":
            return text
        ck = self._cache_key(col, text, entity)
        if ck in self._cache:
            return self._cache[ck]
        out = self._masker(rule, col).mask(text, entity)
        if len(self._cache) < CACHE_LIMIT:
            self._cache[ck] = out
        return out

    # ------------------------------------------------------------------ headers
    @staticmethod
    def _resolve_columns(header: List[Any], rule: FileRule, where: str, res: FileResult):
        """Map configured column names (case/space-insensitive) to header positions."""
        index = {}
        for i, h in enumerate(header):
            if h is not None and str(h).strip():
                index.setdefault(str(h).strip().lower(), i)
        targets = []
        missing = []
        for col in rule.columns:
            pos = index.get(col.name.lower())
            if pos is None:
                missing.append(col.name)
                continue
            ent = None
            if col.entity_column:
                ent = index.get(col.entity_column.lower())
                if ent is None:
                    missing.append("%s (entity_column of %s)" % (col.entity_column, col.name))
                    continue
            targets.append((col, pos, ent))
        if missing:
            msg = "%s: column(s) not found: %s" % (where, ", ".join(missing))
            if rule.on_missing_column == "error":
                # The header row is not echoed: with a wrong header_row (or no header) it would be data.
                raise FileRuleError(msg + " (the header row has %d columns; check header_row and the column "
                                    "names, or run files-validate)" % sum(1 for h in header if h is not None))
            log.warning(msg)
            res.message = (res.message + "; " if res.message else "") + msg
        return targets

    def _apply(self, rule, res, targets, values: List[Any], to_text, from_text, where) -> List[Any]:
        """Mask one row (list of raw values). Entities are read from the ORIGINAL row."""
        # Hot path (runs for every row): cache lookups inline, row copy only when an entity column is used.
        original = list(values) if any(ent is not None for _, _, ent in targets) else values
        cache = self._cache
        n = len(values)
        counts = res.masked_cells
        for col, pos, ent in targets:
            if pos >= n:
                continue
            raw = original[pos]
            text = to_text(raw)
            if not text or text.isspace():
                continue
            entity = to_text(original[ent]) if ent is not None and ent < n else None
            ck = (self._column_cache_id(col), entity if col.entity_column else None, text)
            masked = cache.get(ck, _MISSING)
            if masked is _MISSING:
                try:
                    masked = self.mask_text(rule, col, text, entity)
                except MaskerError as exc:
                    res.value_errors += 1
                    if rule.on_value_error == "fail":
                        raise FileRuleError("%s, column %s: %s" % (where, col.name, exc)) from None
                    if rule.on_value_error == "keep":
                        continue
                    masked = None
            values[pos] = from_text(raw, masked)
            counts[col.name] = counts.get(col.name, 0) + 1
        return values

    # ------------------------------------------------------------------ csv
    def _process_csv(self, src: Path, dst: Optional[Path], rule: FileRule, res: FileResult) -> None:
        opts = rule.csv
        delimiter = "\t" if src.suffix.lower() == ".tsv" and opts.get("delimiter", ",") == "," else opts["delimiter"]
        enc = opts.get("encoding", "utf-8-sig")
        csv.field_size_limit(max(csv.field_size_limit(), int(opts.get("max_field_kb", 1024)) * 1024))
        with src.open("r", encoding=enc, newline="") as fin:
            reader = csv.reader(fin, delimiter=delimiter, quotechar=opts.get("quotechar", '"'))
            header = next(reader, None)
            if header is None:
                raise FileRuleError("%s is empty" % src.name)
            targets = self._resolve_columns(header, rule, src.name, res)
            if targets and self.workers > 1 and src.stat().st_size >= PARALLEL_MIN_BYTES:
                # pass 1: distinct values -> parallel masking; pass 2 below reads the file again
                with src.open("r", encoding=enc, newline="") as pre:
                    pre_reader = csv.reader(pre, delimiter=delimiter, quotechar=opts.get("quotechar", '"'))
                    next(pre_reader, None)
                    n = self.prewarm(rule, targets, pre_reader, lambda v: v)
                log.info("%s: %d distinct values masked by %d worker processes", src.name, n, self.workers)
            fout = dst.open("w", encoding=enc, newline="") if dst else None
            try:
                writer = csv.writer(fout, delimiter=delimiter, quotechar=opts.get("quotechar", '"'),
                                    quoting=csv.QUOTE_MINIMAL) if fout else None
                if writer:
                    writer.writerow(header)
                for line_no, row in enumerate(reader, start=2):
                    out = self._apply(rule, res, targets, row, lambda v: v, lambda o, m: "" if m is None else m,
                                      "%s line %d" % (src.name, line_no))
                    res.rows += 1
                    if writer:
                        writer.writerow(out)
            finally:
                if fout:
                    fout.close()

    # ------------------------------------------------------------------ excel
    @staticmethod
    def _openpyxl():
        try:
            import openpyxl
        except ImportError:
            raise FileRuleError("Excel support needs the openpyxl library. ActiveState: add the package 'openpyxl' "
                                "to your project and run 'state pull' (pip: pip install openpyxl)") from None
        if not getattr(openpyxl, "DEFUSEDXML", False):
            # Excel files are zipped XML; without defusedxml a crafted file could use XML entity expansion.
            raise FileRuleError("Excel support requires the 'defusedxml' package (protects against malicious "
                                "XML in workbooks). ActiveState: add 'defusedxml' and run 'state pull'.")
        return openpyxl

    def _sheet_names(self, wb, rule: FileRule, src: Path) -> List[str]:
        if not rule.sheets:
            return list(wb.sheetnames)
        lower = {n.lower(): n for n in wb.sheetnames}
        missing = [s for s in rule.sheets if s.lower() not in lower]
        if missing:
            msg = "%s: sheet(s) not found: %s (sheets: %s)" % (src.name, ", ".join(missing), ", ".join(wb.sheetnames))
            if rule.on_missing_column == "error":
                raise FileRuleError(msg)
            log.warning(msg)
        return [lower[s.lower()] for s in rule.sheets if s.lower() in lower]

    def _process_excel(self, src: Path, dst: Optional[Path], rule: FileRule, res: FileResult) -> None:
        openpyxl = self._openpyxl()
        hr = int(rule.excel.get("header_row", 1))
        keep_vba = src.suffix.lower() == ".xlsm"
        if rule.excel.get("mode") == "streaming":
            self._excel_streaming(openpyxl, src, dst, rule, res, hr)
            return
        wb = openpyxl.load_workbook(src, keep_vba=keep_vba)
        try:
            for name in self._sheet_names(wb, rule, src):
                ws = wb[name]
                header = [c.value for c in ws[hr]] if ws.max_row >= hr else []
                where = "%s [%s]" % (src.name, name)
                targets = self._resolve_columns(header, rule, where, res) if header else []
                res.sheets.append(name)
                if not targets:
                    continue
                if self.workers > 1 and src.stat().st_size >= PARALLEL_MIN_BYTES:
                    self.prewarm(rule, targets, ([c.value for c in r] for r in ws.iter_rows(min_row=hr + 1)),
                                 cell_to_text)
                for row in ws.iter_rows(min_row=hr + 1):
                    values = [c.value for c in row]
                    formulas = {pos for _, pos, _ in targets if pos < len(row) and row[pos].data_type == "f"}
                    for pos in formulas:  # formulas are left untouched: they are not data
                        values[pos] = None
                        res.formulas_skipped += 1
                    out = self._apply(rule, res, targets, values, cell_to_text, text_to_cell,
                                      "%s row %d" % (where, row[0].row))
                    for _, pos, _ in targets:
                        if pos < len(row) and pos not in formulas and out[pos] != row[pos].value:
                            row[pos].value = out[pos]
                    res.rows += 1
            if dst:
                wb.save(dst)
        finally:
            wb.close()

    def _excel_streaming(self, openpyxl, src, dst, rule, res, hr) -> None:
        """Values only (no formatting) - for very large workbooks."""
        wb_in = openpyxl.load_workbook(src, read_only=True)
        wb_out = openpyxl.Workbook(write_only=True) if dst else None
        try:
            for name in self._sheet_names(wb_in, rule, src):
                ws = wb_in[name]
                ws_out = wb_out.create_sheet(title=name) if wb_out else None
                targets = []
                where = "%s [%s]" % (src.name, name)
                res.sheets.append(name)
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    values = list(row)
                    if r_idx == hr:
                        targets = self._resolve_columns(values, rule, where, res)
                    elif r_idx > hr and targets:
                        values = self._apply(rule, res, targets, values, cell_to_text, text_to_cell,
                                             "%s row %d" % (where, r_idx))
                        res.rows += 1
                    if ws_out is not None:
                        ws_out.append(values)
            if wb_out is not None:
                wb_out.save(dst)
        finally:
            wb_in.close()

    # ------------------------------------------------------------------ one file
    def process(self, src: Path, input_root: Path, output_dir: Path, rule: FileRule, dry_run: bool,
                overwrite: bool) -> FileResult:
        fmt = file_format(src, rule)
        dst = output_path_for(src, input_root, output_dir, rule)
        res = FileResult(str(src), str(dst), rule.name, fmt)
        t0 = time.time()
        size_mb = src.stat().st_size / (1024 * 1024)
        if size_mb > rule.max_file_mb:
            raise FileRuleError("%s is %.0f MB, above max_file_mb=%d in file_rules.yaml" % (src.name, size_mb,
                                                                                        rule.max_file_mb))
        if dst.resolve() == src.resolve():
            raise FileRuleError("%s: output would overwrite the input file; use a different file_output_dir or "
                                "output_suffix" % src)
        if not dry_run and dst.exists() and not overwrite:
            raise FileRuleError("%s already exists (use --overwrite to replace it)" % dst)
        res.input_sha256 = sha256_file(src)
        tmp = None
        if not dry_run:
            dst.parent.mkdir(parents=True, exist_ok=True)
            tmp = dst.with_name(dst.name + ".part")
        try:
            if fmt == "csv":
                self._process_csv(src, tmp, rule, res)
            else:
                self._process_excel(src, tmp, rule, res)
            if tmp is not None:
                os.replace(tmp, dst)  # atomic: a partial output never looks finished
                res.output_sha256 = sha256_file(dst)
            res.status = "DRY_RUN" if dry_run else "COMPLETED"
        finally:
            if tmp is not None and tmp.exists():
                tmp.unlink()
            res.seconds = round(time.time() - t0, 2)
        return res
