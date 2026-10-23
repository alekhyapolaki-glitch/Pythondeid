"""De-identification of CSV and Excel files, driven by config/file_rules.yaml.

The same domains, masking policy and secret key as the database are used, so a
value masks identically in a file and in SQL Server (e.g. an SSN in an extract
gets the same masked SSN as the Customers.SSN column).

file_rules.yaml (location: paths.yaml -> file_rules)
----------------------------------------------------
defaults:                       # apply to every file entry unless overridden there
  output_suffix: _deid          # customers.csv -> customers_deid.csv
  on_missing_column: error      # error | warn  (configured column not in the file/sheet)
  on_value_error: fail          # fail | blank | keep  (value the masker cannot handle, e.g. bad date)
  csv: {delimiter: ",", encoding: utf-8-sig, quotechar: '"'}
  excel: {mode: in_place, header_row: 1}   # in_place keeps formatting | streaming for very large files
files:
  - name: customer_extract
    pattern: "customers*.csv"   # glob, relative to file_input_dir ("**/" for subfolders)
    columns:
      - {name: Email, domain: email}
      - {name: SSN, domain: ssn}
      - {name: DOB, domain: person_dates, entity_column: CustomerId}
  - name: account_workbook
    pattern: "accounts*.xlsx"
    sheets: [Accounts]          # default: every sheet
    columns:
      - {name: Account Number, domain: account_number}

Original values are never written to logs or audit files; the audit records file
fingerprints (SHA-256), row counts and per-column masked-cell counts only.
"""

from __future__ import annotations

import csv
import datetime as _dt
import fnmatch
import hashlib
import json
import logging
import os
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .config import ConfigError, Rules, _read_yaml
from .maskers import BaseMasker, KeyRing, MaskerError, build_maskers, column_masker, get_masker_class

log = logging.getLogger("deid.files")

CSV_EXT = {".csv", ".txt", ".tsv"}
EXCEL_EXT = {".xlsx", ".xlsm"}
DEFAULTS: Dict[str, Any] = {
    "output_suffix": "_deid",
    "on_missing_column": "error",
    "on_value_error": "fail",
    "csv": {"delimiter": ",", "encoding": "utf-8-sig", "quotechar": '"'},
    "excel": {"mode": "in_place", "header_row": 1},
}
CACHE_LIMIT = 500000


class FileRuleError(ConfigError):
    pass


# =============================================================================
# Config
# =============================================================================
@dataclass
class FileColumn:
    name: str
    domain: str
    entity_column: str = ""
    overrides: Dict[str, Any] = field(default_factory=dict)


@dataclass
class FileRule:
    name: str
    pattern: str
    columns: List[FileColumn]
    format: str = ""                     # csv | excel | "" = from the extension
    sheets: List[str] = field(default_factory=list)
    output_suffix: str = "_deid"
    on_missing_column: str = "error"
    on_value_error: str = "fail"
    csv: Dict[str, Any] = field(default_factory=dict)
    excel: Dict[str, Any] = field(default_factory=dict)

    def matches(self, relative: str) -> bool:
        rel = relative.replace("\\", "/")
        pat = self.pattern.replace("\\", "/")
        return fnmatch.fnmatch(rel.lower(), pat.lower()) or fnmatch.fnmatch(Path(rel).name.lower(), pat.lower())


@dataclass
class FileRules:
    files: List[FileRule]
    source_path: Optional[Path] = None
    content_hash: str = ""


def _merge(a: Dict[str, Any], b: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(a)
    for k, v in (b or {}).items():
        out[k] = _merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out


def parse_file_rules(raw: Dict[str, Any], rules: Rules, source: str = "<file_rules>") -> FileRules:
    unknown = set(raw) - {"defaults", "files"}
    if unknown:
        raise FileRuleError("%s: unknown top-level key(s) %s (expected defaults, files)" % (source, sorted(unknown)))
    defaults = _merge(DEFAULTS, raw.get("defaults") or {})
    bad = set(defaults) - set(DEFAULTS)
    if bad:
        raise FileRuleError("%s: unknown defaults key(s) %s" % (source, sorted(bad)))
    out: List[FileRule] = []
    names = set()
    for i, f in enumerate(raw.get("files") or []):
        where = "%s files[%d]" % (source, i)
        if not isinstance(f, dict) or not f.get("pattern"):
            raise FileRuleError("%s: each file entry needs a 'pattern' (e.g. customers*.csv)" % where)
        allowed = {"name", "pattern", "format", "sheets", "columns"} | set(DEFAULTS)
        extra = set(f) - allowed
        if extra:
            raise FileRuleError("%s: unknown key(s) %s. Allowed: %s" % (where, sorted(extra), sorted(allowed)))
        name = str(f.get("name") or f["pattern"])
        if name in names:
            raise FileRuleError("%s: duplicate file entry name '%s'" % (where, name))
        names.add(name)
        opts = _merge(defaults, {k: f[k] for k in DEFAULTS if k in f})
        fmt = str(f.get("format") or "").lower()
        if fmt not in ("", "csv", "excel"):
            raise FileRuleError("%s: format must be csv or excel" % where)
        if opts["on_missing_column"] not in ("error", "warn"):
            raise FileRuleError("%s: on_missing_column must be error or warn" % where)
        if opts["on_value_error"] not in ("fail", "blank", "keep"):
            raise FileRuleError("%s: on_value_error must be fail, blank or keep" % where)
        if opts["excel"].get("mode") not in ("in_place", "streaming"):
            raise FileRuleError("%s: excel.mode must be in_place or streaming" % where)
        if not isinstance(opts["excel"].get("header_row"), int) or opts["excel"]["header_row"] < 1:
            raise FileRuleError("%s: excel.header_row must be a positive integer" % where)
        if len(str(opts["csv"].get("delimiter", ","))) != 1:
            raise FileRuleError("%s: csv.delimiter must be one character (use '\\t' for tab)" % where)
        cols: List[FileColumn] = []
        seen = set()
        for j, c in enumerate(f.get("columns") or []):
            cw = "%s.columns[%d]" % (where, j)
            if not isinstance(c, dict) or not c.get("name") or not c.get("domain"):
                raise FileRuleError("%s: each column needs 'name' (header text) and 'domain'" % cw)
            if str(c["name"]).strip().lower() in seen:
                raise FileRuleError("%s: column '%s' listed twice" % (cw, c["name"]))
            seen.add(str(c["name"]).strip().lower())
            dom = rules.domains.get(str(c["domain"]))
            if dom is None:
                raise FileRuleError("%s: domain '%s' is not defined in %s. File columns use the SAME domains as the "
                                    "database so values mask identically." % (cw, c["domain"], rules.source_path))
            cls = get_masker_class(dom.type)
            overrides = {k: v for k, v in c.items() if k not in ("name", "domain", "entity_column")}
            bad = set(overrides) - set(cls.format_params)
            if bad:
                raise FileRuleError("%s: unknown key(s) %s (format overrides allowed for %s: %s)"
                                    % (cw, sorted(bad), dom.type, sorted(cls.format_params) or "none"))
            needs_entity = cls.entity_required(dom.params)
            if needs_entity and not c.get("entity_column"):
                raise FileRuleError("%s: domain '%s' shifts dates per person; add entity_column (the person id "
                                    "column in this file)" % (cw, dom.name))
            cols.append(FileColumn(str(c["name"]).strip(), dom.name, str(c.get("entity_column") or "").strip(),
                                   overrides))
        if not cols:
            raise FileRuleError("%s: no columns configured" % where)
        out.append(FileRule(name, str(f["pattern"]), cols, fmt, [str(s) for s in (f.get("sheets") or [])],
                            str(opts["output_suffix"]), opts["on_missing_column"], opts["on_value_error"],
                            dict(opts["csv"]), dict(opts["excel"])))
    blob = json.dumps(raw, sort_keys=True, default=str).encode("utf-8")
    return FileRules(out, content_hash=hashlib.sha256(blob).hexdigest()[:16])


def load_file_rules(path: str, rules: Rules) -> FileRules:
    p = Path(path)
    if not p.exists():
        raise FileRuleError("File rules not found: %s. Create one with: python main.py files-discover" % p)
    fr = parse_file_rules(_read_yaml(p), rules, str(p))
    fr.source_path = p
    if not fr.files:
        raise FileRuleError("%s defines no files" % p)
    return fr


def file_format(path: Path, rule: Optional[FileRule] = None) -> str:
    if rule is not None and rule.format:
        return rule.format
    ext = path.suffix.lower()
    if ext in CSV_EXT:
        return "csv"
    if ext in EXCEL_EXT:
        return "excel"
    if ext == ".xls":
        raise FileRuleError("%s: old .xls format is not supported; save it as .xlsx" % path.name)
    raise FileRuleError("%s: unsupported file type %s (csv, txt, tsv, xlsx, xlsm)" % (path.name, ext))


def find_inputs(input_path: str, file_rules: FileRules) -> Tuple[List[Tuple[Path, FileRule]], List[Path]]:
    """Returns ([(file, rule)], [files with no matching rule])."""
    root = Path(input_path)
    if root.is_file():
        candidates, base = [root], root.parent
    elif root.is_dir():
        candidates = sorted(p for p in root.rglob("*") if p.is_file() and not p.name.startswith(("~$", ".")))
        base = root
    else:
        raise FileRuleError("Input path not found: %s (set file_input_dir in paths.yaml or pass --input)" % root)
    matched, unmatched = [], []
    for p in candidates:
        rel = str(p.relative_to(base))
        rule = next((r for r in file_rules.files if r.matches(rel)), None)
        (matched if rule else unmatched).append((p, rule) if rule else p)
    return matched, unmatched


def output_path_for(src: Path, input_root: Path, output_dir: Path, rule: FileRule) -> Path:
    base = input_root if input_root.is_dir() else input_root.parent
    rel = src.relative_to(base)
    return output_dir / rel.parent / ("%s%s%s" % (src.stem, rule.output_suffix, src.suffix))


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# =============================================================================
# Value conversion (Excel cell types <-> text the maskers work on)
# =============================================================================
_INT_RE = re.compile(r"^-?\d+$")


def cell_to_text(value: Any) -> Optional[str]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, _dt.datetime):
        return value.isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def text_to_cell(original: Any, masked: Optional[str]) -> Any:
    """Write back with the original cell type where possible (numbers stay numbers, dates stay dates)."""
    if masked is None:
        return None
    if isinstance(original, _dt.datetime):
        try:
            return _dt.datetime.fromisoformat(masked)
        except ValueError:
            return masked
    if isinstance(original, _dt.date):
        try:
            return _dt.date.fromisoformat(masked[:10])
        except ValueError:
            return masked
    if isinstance(original, (int, float)) and not isinstance(original, bool):
        if _INT_RE.match(masked or ""):
            return int(masked)
        try:
            return float(masked)
        except ValueError:
            return masked
    return masked


# =============================================================================
# Processing
# =============================================================================
@dataclass
class FileResult:
    input_path: str
    output_path: str
    rule: str
    format: str
    status: str = "PENDING"
    rows: int = 0
    sheets: List[str] = field(default_factory=list)
    masked_cells: Dict[str, int] = field(default_factory=dict)
    value_errors: int = 0
    formulas_skipped: int = 0
    input_sha256: str = ""
    output_sha256: str = ""
    seconds: float = 0.0
    message: str = ""


class FileDeidentifier:
    def __init__(self, rules: Rules, file_rules: FileRules, keyring: KeyRing):
        self.rules = rules
        self.file_rules = file_rules
        self.keyring = keyring
        self.maskers = build_maskers(rules.domains.values(), keyring)
        self._col_maskers: Dict[Tuple[str, str], BaseMasker] = {}
        self._cache: Dict[Tuple[str, str, Optional[str], str], Optional[str]] = {}

    def _masker(self, rule: FileRule, col: FileColumn) -> BaseMasker:
        key = (rule.name, col.name)
        m = self._col_maskers.get(key)
        if m is None:
            m = column_masker(self.maskers[col.domain], self.keyring, col.overrides)
            self._col_maskers[key] = m
        return m

    def mask_text(self, rule: FileRule, col: FileColumn, text: Optional[str], entity: Optional[str]) -> Optional[str]:
        if text is None or text.strip() == "":
            return text
        ck = (col.domain, json.dumps(col.overrides, sort_keys=True), entity if col.entity_column else None, text)
        if ck in self._cache:
            return self._cache[ck]
        out = self._masker(rule, col).mask(text, entity)
        if len(self._cache) < CACHE_LIMIT:
            self._cache[ck] = out
        return out

    # ------------------------------------------------------------------ headers
    @staticmethod
    def _resolve_columns(header: List[Any], rule: FileRule, where: str, res: FileResult):
        """Map configured column names (case/space-insensitive) to header positions."""
        index = {}
        for i, h in enumerate(header):
            if h is not None and str(h).strip():
                index.setdefault(str(h).strip().lower(), i)
        targets = []
        missing = []
        for col in rule.columns:
            pos = index.get(col.name.lower())
            if pos is None:
                missing.append(col.name)
                continue
            ent = None
            if col.entity_column:
                ent = index.get(col.entity_column.lower())
                if ent is None:
                    missing.append("%s (entity_column of %s)" % (col.entity_column, col.name))
                    continue
            targets.append((col, pos, ent))
        if missing:
            msg = "%s: column(s) not found: %s" % (where, ", ".join(missing))
            if rule.on_missing_column == "error":
                raise FileRuleError(msg + ". Header is: %s" % ", ".join(str(h) for h in header if h is not None))
            log.warning(msg)
            res.message = (res.message + "; " if res.message else "") + msg
        return targets

    def _apply(self, rule, res, targets, values: List[Any], to_text, from_text, where) -> List[Any]:
        """Mask one row (list of raw values). Entities are read from the ORIGINAL row."""
        original = list(values)
        for col, pos, ent in targets:
            if pos >= len(values):
                continue
            raw = original[pos]
            text = to_text(raw)
            if text is None or text.strip() == "":
                continue
            entity = to_text(original[ent]) if ent is not None and ent < len(original) else None
            try:
                masked = self.mask_text(rule, col, text, entity)
            except MaskerError as exc:
                res.value_errors += 1
                if rule.on_value_error == "fail":
                    raise FileRuleError("%s, column %s: %s" % (where, col.name, exc))
                masked = None if rule.on_value_error == "blank" else text
                if rule.on_value_error == "keep":
                    continue
            values[pos] = from_text(raw, masked)
            res.masked_cells[col.name] = res.masked_cells.get(col.name, 0) + 1
        return values

    # ------------------------------------------------------------------ csv
    def _process_csv(self, src: Path, dst: Optional[Path], rule: FileRule, res: FileResult) -> None:
        opts = rule.csv
        delimiter = "\t" if src.suffix.lower() == ".tsv" and opts.get("delimiter", ",") == "," else opts["delimiter"]
        enc = opts.get("encoding", "utf-8-sig")
        with src.open("r", encoding=enc, newline="") as fin:
            reader = csv.reader(fin, delimiter=delimiter, quotechar=opts.get("quotechar", '"'))
            header = next(reader, None)
            if header is None:
                raise FileRuleError("%s is empty" % src.name)
            targets = self._resolve_columns(header, rule, src.name, res)
            fout = dst.open("w", encoding=enc, newline="") if dst else None
            try:
                writer = csv.writer(fout, delimiter=delimiter, quotechar=opts.get("quotechar", '"'),
                                    quoting=csv.QUOTE_MINIMAL) if fout else None
                if writer:
                    writer.writerow(header)
                for line_no, row in enumerate(reader, start=2):
                    out = self._apply(rule, res, targets, row, lambda v: v, lambda o, m: "" if m is None else m,
                                      "%s line %d" % (src.name, line_no))
                    res.rows += 1
                    if writer:
                        writer.writerow(out)
            finally:
                if fout:
                    fout.close()

    # ------------------------------------------------------------------ excel
    @staticmethod
    def _openpyxl():
        try:
            import openpyxl  # noqa: WPS433
        except ImportError:
            raise FileRuleError("Excel support needs the openpyxl library. ActiveState: add the package 'openpyxl' "
                                "to your project and run 'state pull' (pip: pip install openpyxl)")
        return openpyxl

    def _sheet_names(self, wb, rule: FileRule, src: Path) -> List[str]:
        if not rule.sheets:
            return list(wb.sheetnames)
        lower = {n.lower(): n for n in wb.sheetnames}
        missing = [s for s in rule.sheets if s.lower() not in lower]
        if missing:
            msg = "%s: sheet(s) not found: %s (sheets: %s)" % (src.name, ", ".join(missing), ", ".join(wb.sheetnames))
            if rule.on_missing_column == "error":
                raise FileRuleError(msg)
            log.warning(msg)
        return [lower[s.lower()] for s in rule.sheets if s.lower() in lower]

    def _process_excel(self, src: Path, dst: Optional[Path], rule: FileRule, res: FileResult) -> None:
        openpyxl = self._openpyxl()
        hr = int(rule.excel.get("header_row", 1))
        keep_vba = src.suffix.lower() == ".xlsm"
        if rule.excel.get("mode") == "streaming":
            self._excel_streaming(openpyxl, src, dst, rule, res, hr)
            return
        wb = openpyxl.load_workbook(src, keep_vba=keep_vba)
        try:
            for name in self._sheet_names(wb, rule, src):
                ws = wb[name]
                header = [c.value for c in ws[hr]] if ws.max_row >= hr else []
                where = "%s [%s]" % (src.name, name)
                targets = self._resolve_columns(header, rule, where, res) if header else []
                res.sheets.append(name)
                if not targets:
                    continue
                for row in ws.iter_rows(min_row=hr + 1):
                    values = [c.value for c in row]
                    formulas = {pos for _, pos, _ in targets if pos < len(row) and row[pos].data_type == "f"}
                    for pos in formulas:  # formulas are left untouched: they are not data
                        values[pos] = None
                        res.formulas_skipped += 1
                    out = self._apply(rule, res, targets, values, cell_to_text, text_to_cell,
                                      "%s row %d" % (where, row[0].row))
                    for _, pos, _ in targets:
                        if pos < len(row) and pos not in formulas and out[pos] != row[pos].value:
                            row[pos].value = out[pos]
                    res.rows += 1
            if dst:
                wb.save(dst)
        finally:
            wb.close()

    def _excel_streaming(self, openpyxl, src, dst, rule, res, hr) -> None:
        """Values only (no formatting) - for very large workbooks."""
        wb_in = openpyxl.load_workbook(src, read_only=True)
        wb_out = openpyxl.Workbook(write_only=True) if dst else None
        try:
            for name in self._sheet_names(wb_in, rule, src):
                ws = wb_in[name]
                ws_out = wb_out.create_sheet(title=name) if wb_out else None
                targets = []
                where = "%s [%s]" % (src.name, name)
                res.sheets.append(name)
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    values = list(row)
                    if r_idx == hr:
                        targets = self._resolve_columns(values, rule, where, res)
                    elif r_idx > hr and targets:
                        values = self._apply(rule, res, targets, values, cell_to_text, text_to_cell,
                                             "%s row %d" % (where, r_idx))
                        res.rows += 1
                    if ws_out is not None:
                        ws_out.append(values)
            if wb_out is not None:
                wb_out.save(dst)
        finally:
            wb_in.close()

    # ------------------------------------------------------------------ one file
    def process(self, src: Path, input_root: Path, output_dir: Path, rule: FileRule, dry_run: bool,
                overwrite: bool) -> FileResult:
        fmt = file_format(src, rule)
        dst = output_path_for(src, input_root, output_dir, rule)
        res = FileResult(str(src), str(dst), rule.name, fmt)
        t0 = time.time()
        if dst.resolve() == src.resolve():
            raise FileRuleError("%s: output would overwrite the input file; use a different file_output_dir or "
                                "output_suffix" % src)
        if not dry_run and dst.exists() and not overwrite:
            raise FileRuleError("%s already exists (use --overwrite to replace it)" % dst)
        res.input_sha256 = sha256_file(src)
        tmp = None
        if not dry_run:
            dst.parent.mkdir(parents=True, exist_ok=True)
            tmp = dst.with_name(dst.name + ".part")
        try:
            if fmt == "csv":
                self._process_csv(src, tmp, rule, res)
            else:
                self._process_excel(src, tmp, rule, res)
            if tmp is not None:
                os.replace(tmp, dst)  # atomic: a partial output never looks finished
                res.output_sha256 = sha256_file(dst)
            res.status = "DRY_RUN" if dry_run else "COMPLETED"
        finally:
            if tmp is not None and tmp.exists():
                tmp.unlink()
            res.seconds = round(time.time() - t0, 2)
        return res


# =============================================================================
# Discovery: draft a file_rules.yaml from file headers
# =============================================================================
def read_headers(path: Path, header_row: int = 1, csv_opts: Optional[Dict[str, Any]] = None) -> Dict[str, List[str]]:
    """{sheet name or '' for CSV: [header cells]}"""
    fmt = file_format(path)
    if fmt == "csv":
        opts = dict(DEFAULTS["csv"], **(csv_opts or {}))
        delim = "\t" if path.suffix.lower() == ".tsv" else opts["delimiter"]
        with path.open("r", encoding=opts["encoding"], newline="") as fh:
            row = next(csv.reader(fh, delimiter=delim, quotechar=opts["quotechar"]), [])
        return {"": [h.strip() for h in row]}
    openpyxl = FileDeidentifier._openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        out = {}
        for name in wb.sheetnames:
            rows = wb[name].iter_rows(min_row=header_row, max_row=header_row, values_only=True)
            first = next(rows, ())
            out[name] = [str(h).strip() for h in first if h is not None and str(h).strip()]
        return out
    finally:
        wb.close()


def draft_file_rules(input_dir: str, rules: Rules) -> Tuple[str, int]:
    """YAML text of a draft file_rules.yaml with suggested domains (uses the DB column-name heuristics)."""
    from .discovery import suggest_type
    from .schema import ColumnInfo
    by_type: Dict[str, str] = {}
    for d in rules.domains.values():
        by_type.setdefault(d.type, d.name)
    root = Path(input_dir)
    files = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in CSV_EXT | EXCEL_EXT
                   and not p.name.startswith(("~$", ".")))
    lines = ["# Draft generated by: python main.py files-discover  - review before use.",
             "# Each column's domain must exist in the rules file; same domain = same masked value as the database.",
             "# Suggested domains are guesses from header names; delete columns that are not PII.",
             "defaults:",
             "  output_suffix: _deid",
             "  on_missing_column: error      # error | warn",
             "  on_value_error: fail          # fail | blank | keep",
             "  csv: {delimiter: \",\", encoding: utf-8-sig, quotechar: '\"'}",
             "  excel: {mode: in_place, header_row: 1}   # in_place keeps formatting | streaming for huge files",
             "files:"]
    count = 0
    for f in files:
        try:
            headers = read_headers(f)
        except (FileRuleError, OSError, UnicodeDecodeError) as exc:
            lines.append("  # skipped %s: %s" % (f.name, exc))
            continue
        rel = str(f.relative_to(root)).replace("\\", "/")
        sheets = [s for s in headers if s]
        entry = ["  - name: %s" % re.sub(r"[^A-Za-z0-9_]+", "_", f.stem).strip("_").lower(),
                 "    pattern: \"%s\"" % rel]
        if sheets:
            entry.append("    sheets: [%s]" % ", ".join(json.dumps(s) for s in sheets))
        entry.append("    columns:")
        seen = set()
        any_col = False
        for sheet, cols in headers.items():
            for h in cols:
                if h.lower() in seen:
                    continue
                seen.add(h.lower())
                mtype = suggest_type(ColumnInfo("", "", h, "varchar", 100, 0, 0, True))
                if mtype:
                    domain = by_type.get(mtype, mtype)
                    entry.append("      - {name: %s, domain: %s}" % (json.dumps(h), domain))
                    any_col = True
                else:
                    entry.append("      # - {name: %s, domain: ???}" % json.dumps(h))
        if not any_col:  # nothing PII-looking: keep the entry commented out for review
            entry = ["  # no PII-looking headers found in %s; uncomment and add columns if needed" % f.name] + \
                ["  # " + e.lstrip() if not e.lstrip().startswith("#") else "  " + e.lstrip() for e in entry]
        lines.extend(entry)
        count += 1
    if not count:
        lines.append("  []  # no .csv/.txt/.tsv/.xlsx/.xlsm files found in %s" % root)
    return "\n".join(lines) + "\n", count
